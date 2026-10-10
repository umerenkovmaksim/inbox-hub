import json
from pathlib import Path

import docx
import openpyxl
import pptx
from pypdf import PdfReader


def docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def pptx_text(path: Path) -> str:
    presentation = pptx.Presentation(str(path))
    parts = []
    for number, slide in enumerate(presentation.slides, 1):
        parts.append(f"--- slide {number}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def xlsx_text(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"--- sheet {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append("\t".join("" if v is None else str(v) for v in row))
    workbook.close()
    return "\n".join(parts)


def pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def ipynb_text(path: Path) -> str:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    parts = []
    for cell in notebook.get("cells", []):
        source = cell.get("source", "")
        source = "".join(source) if isinstance(source, list) else source
        if cell.get("cell_type") == "code":
            parts.append(f"```python\n{source}\n```")
        else:
            parts.append(source)
    return "\n\n".join(parts)


def plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")
