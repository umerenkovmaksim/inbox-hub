import json
import zipfile

import docx
import openpyxl
import pptx
from pypdf import PdfWriter

from university_pz_parser.extractors.registry import ExtractorRegistry


def _pdf_with_text(path, text):
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET".encode()
    objs[3] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer << /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objs) + 1,
        xref,
    )
    path.write_bytes(bytes(out))


def test_docx(tmp_path):
    d = docx.Document()
    d.add_paragraph("Практическое задание 1")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Вариант"
    t.cell(0, 1).text = "3"
    p = tmp_path / "a.docx"
    d.save(p)
    r = ExtractorRegistry().extract(p, "a.docx", None)
    assert r.status == "ok"
    assert "Практическое задание 1" in r.text and "Вариант\t3" in r.text


def test_pptx_xlsx_ipynb(tmp_path):
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Лекция 3"
    prs.save(tmp_path / "l.pptx")
    wb = openpyxl.Workbook()
    wb.active.append(["ФИО", "Балл"])
    wb.save(tmp_path / "t.xlsx")
    nb = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Занятие 1"]},
            {"cell_type": "code", "source": "print(1)"},
        ]
    }
    (tmp_path / "n.ipynb").write_text(json.dumps(nb))
    reg = ExtractorRegistry()
    assert "Лекция 3" in reg.extract(tmp_path / "l.pptx", "l.pptx", None).text
    assert "ФИО\tБалл" in reg.extract(tmp_path / "t.xlsx", None, None).text
    nbt = reg.extract(tmp_path / "n.ipynb", None, None).text
    assert "# Занятие 1" in nbt and "print(1)" in nbt


def test_pdf_text_and_scan(tmp_path):
    _pdf_with_text(tmp_path / "t.pdf", "Hello PZ")
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.write(tmp_path / "scan.pdf")
    reg = ExtractorRegistry()
    assert "Hello PZ" in reg.extract(tmp_path / "t.pdf", None, None).text
    assert reg.extract(tmp_path / "scan.pdf", None, None).status == "needs_ocr"


def test_zip_recursive_and_unsupported_and_failed(tmp_path):
    (tmp_path / "in.txt").write_text("внутри архива")
    with zipfile.ZipFile(tmp_path / "a.zip", "w") as z:
        z.write(tmp_path / "in.txt", "dir/in.txt")
    (tmp_path / "x.bin").write_bytes(b"\x00")
    (tmp_path / "bad.docx").write_bytes(b"not a zip")
    reg = ExtractorRegistry()
    zr = reg.extract(tmp_path / "a.zip", None, None)
    assert (
        zr.status == "ok" and "### dir/in.txt" in zr.text and "внутри архива" in zr.text
    )
    assert reg.extract(tmp_path / "x.bin", None, None).status == "unsupported"
    assert reg.extract(tmp_path / "bad.docx", None, None).status == "failed"


def test_truncation(tmp_path):
    (tmp_path / "big.txt").write_text("я" * 50)
    r = ExtractorRegistry(limit=10).extract(tmp_path / "big.txt", None, None)
    assert r.status == "truncated" and len(r.text) == 10
